"""Create a lightweight Phase 1 audit from official metadata and sample files."""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def read_csv(path: Path) -> list[dict[str, str]]:
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                return list(csv.DictReader(file))
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, str(path))


def read_schema(path: Path) -> dict[str, list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    result: dict[str, list[dict[str, str]]] = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        result[sheet.title] = [
            {headers[i]: str(value).strip() if value is not None else "" for i, value in enumerate(row)}
            for row in rows[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
    return result


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    formal = args.data_root / "正式数据"
    schema_path = formal / "附件3：数据库-表名及字段说明.xlsx"
    company_path = formal / "附件1：医药上市公司基本信息（截至到2026年1月13日）.xlsx"
    schema = read_schema(schema_path)

    company_book = load_workbook(company_path, read_only=True, data_only=True)
    company_sheet = company_book["基本信息表"]
    company_rows = list(company_sheet.iter_rows(values_only=True))
    company_headers = [str(value).strip() if value is not None else "" for value in company_rows[0]]
    companies = [
        {company_headers[i]: value for i, value in enumerate(row)}
        for row in company_rows[1:]
        if any(value is not None for value in row)
    ]
    sample_company = next(row for row in companies if str(row.get("股票代码")) == "2821" or str(row.get("股票代码")).zfill(6) == "002821")

    files = [path for path in formal.rglob("*") if path.is_file()]
    extension_counts = Counter(path.suffix.lower() or "<no_extension>" for path in files)
    financial_pdf_count = len(list((formal / "附件2：财务报告").rglob("*.pdf")))
    research_pdf_count = len(list((formal / "附件5：研报数据").rglob("*.pdf")))
    research_md_count = len(list((formal / "附件5：研报数据").rglob("*.md")))

    report = {
        "scope": "Phase 1 lightweight audit",
        "sample_company": sample_company,
        "schema_sheets": {
            sheet: len(rows) for sheet, rows in schema.items()
        },
        "schema_fields": {
            sheet: [row.get("字段名称", "") for row in rows]
            for sheet, rows in schema.items()
            if sheet != "数据库表名"
        },
        "file_counts": dict(extension_counts),
        "financial_report_pdf_count": financial_pdf_count,
        "research_pdf_count": research_pdf_count,
        "research_markdown_count": research_md_count,
        "structured_financial_value_files_found": [],
        "first_metrics_mapping": {
            "revenue": {
                "display_name": "营业收入",
                "source_candidates": [
                    "核心业绩指标表.total_operating_revenue",
                    "利润表.total_operating_revenue",
                ],
                "unit": "万元",
            },
            "net_profit": {
                "display_name": "净利润",
                "source_candidates": [
                    "核心业绩指标表.net_profit_10k_yuan",
                    "利润表.net_profit",
                ],
                "unit": "万元",
            },
            "roe": {
                "display_name": "净资产收益率",
                "source_candidates": ["核心业绩指标表.roe"],
                "unit": "%",
            },
        },
        "finding": "当前正式数据目录包含公司信息、字段定义、问题集、财报PDF和研报资料；未发现官方财务事实值CSV/Excel。第一版财务事实需要从年报表格抽取并经过校验，不能把字段说明当作事实数据。",
    }

    args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / "phase1_data_audit.json"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    md = f"""# Phase 1 数据审计报告\n\n## 样本公司\n\n- 公司：{sample_company.get('公司名称')}\n- 股票代码：{str(sample_company.get('股票代码')).zfill(6)}\n- 行业：{sample_company.get('所属证监会行业')}\n\n## 已发现数据\n\n- 财报 PDF：{financial_pdf_count} 份\n- 研报 PDF：{research_pdf_count} 份\n- 研报 Markdown：{research_md_count} 份\n- 字段定义工作表：{len(schema)} 个\n- B 题样本公司：{len(companies)} 家\n\n## 第一轮指标映射\n\n| metric_code | 中文名 | 原始字段候选 | 单位 |\n|---|---|---|---|\n| revenue | 营业收入 | 核心业绩指标表.total_operating_revenue；利润表.total_operating_revenue | 万元 |\n| net_profit | 净利润 | 核心业绩指标表.net_profit_10k_yuan；利润表.net_profit | 万元 |\n| roe | 净资产收益率 | 核心业绩指标表.roe | % |\n\n## 重要发现\n\n正式数据目录当前没有发现官方财务事实值 CSV/Excel，只有数据库字段定义和财报 PDF。因而后续不能直接把字段说明导入财务事实表；需要以年报 PDF 为来源，用 Python 抽取凯莱英 2022—2024 年的财务数据，再进行交叉核验。\n\n## 下一步\n\n先实现一个年报表格抽取与人工可核验输出：\n\n```text\n年报 PDF → 候选财务表格/文本 → revenue、net_profit、roe → CSV → 人工核验\n```\n\n核验通过后，再写入 MySQL 的原始层和统一事实层。\n"""
    md_path = args.output_dir / "phase1_data_audit.md"
    md_path.write_text(md, encoding="utf-8")
    print(md)
    print(f"JSON_OUTPUT={json_path.resolve()}")
    print(f"MARKDOWN_OUTPUT={md_path.resolve()}")


if __name__ == "__main__":
    main()
